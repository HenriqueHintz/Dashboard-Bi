import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from pathlib import Path

wb = openpyxl.Workbook()

# Cores e Estilos
FONT_TITLE = Font(name="Calibri", size=14, bold=True, color="107C41")
FONT_HEADER = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
FONT_BOLD = Font(name="Calibri", size=10, bold=True, color="000000")
FONT_NORMAL = Font(name="Calibri", size=10, color="000000")
FONT_MUTED = Font(name="Calibri", size=9, color="595959", italic=True)

FILL_HEADER = PatternFill(start_color="107C41", end_color="107C41", fill_type="solid")
FILL_ZEBRA = PatternFill(start_color="F2F9F4", end_color="F2F9F4", fill_type="solid")
FILL_TOTAL = PatternFill(start_color="E2F0D9", end_color="E2F0D9", fill_type="solid")
FILL_BLUE = PatternFill(start_color="E8F1FC", end_color="E8F1FC", fill_type="solid")
FILL_AMBER = PatternFill(start_color="FFF4E5", end_color="FFF4E5", fill_type="solid")

BORDER_THIN = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)
BORDER_DOUBLE_BOTTOM = Border(
    top=Side(style="thin", color="000000"),
    bottom=Side(style="double", color="000000"),
)

# ── ABA 1: ORÇAMENTO MENSAL 50-30-20 ──
ws1 = wb.active
ws1.title = "1. Orçamento Mensal"
ws1.views.sheetView[0].showGridLines = True

ws1["B2"] = "FINQUEST — PLANILHA DE GESTÃO FINANCEIRA (REGRA 50/30/20)"
ws1["B2"].font = FONT_TITLE

ws1["B4"] = "Renda Líquida Mensal:"
ws1["B4"].font = FONT_BOLD
ws1["C4"] = 5500.00
ws1["C4"].font = FONT_TITLE
ws1["C4"].number_format = "R$ #,##0.00"

headers1 = ["#", "Categoria", "Descrição / Item", "Previsto (R$)", "Realizado (R$)", "Diferença (R$)", "% Renda"]
for col_idx, h in enumerate(headers1, start=2):
    cell = ws1.cell(row=6, column=col_idx, value=h)
    cell.font = FONT_HEADER
    cell.fill = FILL_HEADER
    cell.alignment = Alignment(horizontal="center" if col_idx in (2, 3) else "left", vertical="center")

items = [
    # Necessidades
    ("NECESSIDADE", "Aluguel / Condomínio", 1400.00, 1400.00),
    ("NECESSIDADE", "Supermercado & Alimentação", 800.00, 780.00),
    ("NECESSIDADE", "Energia Elétrica & Água", 220.00, 245.00),
    ("NECESSIDADE", "Internet Fibra", 120.00, 120.00),
    ("NECESSIDADE", "Transporte / Combustível", 200.00, 190.00),
    # Desejos
    ("DESEJO", "Lazer & Restaurantes", 550.00, 610.00),
    ("DESEJO", "Streaming & Assinaturas", 95.00, 95.00),
    ("DESEJO", "Academia & Cuidados", 150.00, 150.00),
    ("DESEJO", "Compras Pessoais & Roupas", 300.00, 250.00),
    ("DESEJO", "iFood / Delivery", 250.00, 280.00),
    # Investimentos
    ("INVESTIMENTO", "Reserva de Emergência (CDB 100% CDI)", 600.00, 600.00),
    ("INVESTIMENTO", "Tesouro Direto IPCA+ (Longo Prazo)", 350.00, 350.00),
    ("INVESTIMENTO", "Ações / FIIs (Dividendos)", 200.00, 200.00),
]

start_row = 7
for idx, (cat, name, prev, real) in enumerate(items, start=1):
    r = start_row + idx - 1
    ws1.cell(row=r, column=2, value=idx).alignment = Alignment(horizontal="center")
    ws1.cell(row=r, column=3, value=cat)
    ws1.cell(row=r, column=4, value=name)
    
    c_prev = ws1.cell(row=r, column=5, value=prev)
    c_prev.number_format = "R$ #,##0.00"
    
    c_real = ws1.cell(row=r, column=6, value=real)
    c_real.number_format = "R$ #,##0.00"
    
    # Fórmulas
    c_diff = ws1.cell(row=r, column=7, value=f"=E{r}-F{r}")
    c_diff.number_format = "R$ #,##0.00"
    
    c_pct = ws1.cell(row=r, column=8, value=f"=F{r}/$C$4")
    c_pct.number_format = "0.0%"

    for col in range(2, 9):
        ws1.cell(row=r, column=col).border = BORDER_THIN

end_row = start_row + len(items) - 1
total_row = end_row + 1

ws1.cell(row=total_row, column=2, value="Σ").alignment = Alignment(horizontal="center")
ws1.cell(row=total_row, column=3, value="TOTAIS").font = FONT_BOLD
ws1.cell(row=total_row, column=4, value="TOTAL GERAL DE DESPESAS").font = FONT_BOLD
c_tot_prev = ws1.cell(row=total_row, column=5, value=f"=SUM(E{start_row}:E{end_row})")
c_tot_prev.number_format = "R$ #,##0.00"
c_tot_prev.font = FONT_BOLD

c_tot_real = ws1.cell(row=total_row, column=6, value=f"=SUM(F{start_row}:F{end_row})")
c_tot_real.number_format = "R$ #,##0.00"
c_tot_real.font = FONT_BOLD

c_tot_diff = ws1.cell(row=total_row, column=7, value=f"=E{total_row}-F{total_row}")
c_tot_diff.number_format = "R$ #,##0.00"
c_tot_diff.font = FONT_BOLD

c_tot_pct = ws1.cell(row=total_row, column=8, value=f"=F{total_row}/$C$4")
c_tot_pct.number_format = "0.0%"
c_tot_pct.font = FONT_BOLD

for col in range(2, 9):
    ws1.cell(row=total_row, column=col).fill = FILL_TOTAL
    ws1.cell(row=total_row, column=col).border = BORDER_DOUBLE_BOTTOM

# Resumo da Regra 50/30/20
r_rule = total_row + 3
ws1.cell(row=r_rule, column=2, value="ANÁLISE DA REGRA 50/30/20").font = FONT_BOLD

ws1.cell(row=r_rule+1, column=2, value="1. Necessidades (Meta: ≤50%)")
ws1.cell(row=r_rule+1, column=5, value=f'=SUMIF(C{start_row}:C{end_row}, "NECESSIDADE", F{start_row}:F{end_row})').number_format = "R$ #,##0.00"
ws1.cell(row=r_rule+1, column=6, value=f'=E{r_rule+1}/$C$4').number_format = "0.0%"

ws1.cell(row=r_rule+2, column=2, value="2. Desejos Pessoais (Meta: ≤30%)")
ws1.cell(row=r_rule+2, column=5, value=f'=SUMIF(C{start_row}:C{end_row}, "DESEJO", F{start_row}:F{end_row})').number_format = "R$ #,##0.00"
ws1.cell(row=r_rule+2, column=6, value=f'=E{r_rule+2}/$C$4').number_format = "0.0%"

ws1.cell(row=r_rule+3, column=2, value="3. Investimentos (Meta: ≥20%)")
ws1.cell(row=r_rule+3, column=5, value=f'=SUMIF(C{start_row}:C{end_row}, "INVESTIMENTO", F{start_row}:F{end_row})').number_format = "R$ #,##0.00"
ws1.cell(row=r_rule+3, column=6, value=f'=E{r_rule+3}/$C$4').number_format = "0.0%"

# Ajuste de largura das colunas
for col in ws1.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws1.column_dimensions[col_letter].width = max(max_len + 3, 12)

# Salvar
root_dir = Path(__file__).resolve().parent.parent
out_file = root_dir / "public" / "Planilha_Financeira_Oficial.xlsx"
out_file.parent.mkdir(parents=True, exist_ok=True)
wb.save(str(out_file))

# Salva cópia na raiz do projeto
wb.save(str(root_dir / "Planilha_Financeira_Oficial.xlsx"))
print("Planilha física .xlsx gerada com sucesso!")
